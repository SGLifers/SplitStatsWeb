from __future__ import annotations

import csv
import json
import re
import unicodedata
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / "outputs" / "portalwars2-log-extract" / "Splitgate_match_stats_and_medals.xlsx"
OUT = Path(__file__).resolve().parent / "data.js"
SCOREBOARD = Path(__file__).resolve().parent / "scoreboards.csv"


def encode(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return value.total_seconds()
    return value


def rows(sheet, start_row):
    headers = [cell.value for cell in sheet[start_row]]
    output = []
    for values in sheet.iter_rows(min_row=start_row + 1, values_only=True):
        if not any(value is not None for value in values):
            continue
        output.append({header: encode(value) for header, value in zip(headers, values)})
    return output


def player_key(value):
    key = re.sub(r"[^a-z0-9]", "", unicodedata.normalize("NFKD", str(value)).lower())
    if key.startswith("j0gabonito"):
        return "j0gabonito"
    aliases = {
        "robos729tv": "robos729ttv",
        "luluijn": "lulujin",
        "wgrid93isglifers": "wgrid93",
        "wgrid93sglifers": "wgrid93",
    }
    return aliases.get(key, key)


def scoreboard_rows(path, team_rows):
    team_by_player = {}
    for team in team_rows:
        for column in ("Captain", "Player 2", "Player 3", "Player 4"):
            team_by_player[(int(team["Match"]), player_key(team[column]))] = team["Team"]

    with path.open(newline="", encoding="utf-8-sig") as handle:
        output = []
        for row in csv.DictReader(handle):
            if not row.get("Match number"):
                continue
            match_number = int(row["Match number"])
            output.append({
                "Match": match_number,
                "Team": team_by_player.get((match_number, player_key(row["Player"])), "Unassigned"),
                "Player": row["Player"],
                "XP": int(row["XP"]),
                "Kills": int(row["Kills"]),
                "Damage": int(row["Damage"]),
                "Deaths": int(row["Deaths"]),
                "Hill Time": int(row["Hill Time"]),
            })
        return output


book = load_workbook(BOOK, data_only=True, read_only=True)
matches = rows(book["Match Summary"], 8)[:6]
teams = rows(book["Teams"], 5)
named_winner_overrides = {"1": "Delta", "2": "Alpha", "3": "Charlie", "4": "Delta", "5": "Alpha", "6": "Delta"}
for match in matches:
    override = named_winner_overrides.get(str(match.get("Match")))
    if override:
        match["Winner"] = override
data = {
    "matches": matches,
    "teams": teams,
    "players": rows(book["Player Medals"], 5),
    "events": rows(book["Medal Events"], 5),
    "scoreboards": scoreboard_rows(SCOREBOARD, teams),
    "generatedAt": datetime.now().astimezone().isoformat(),
    "source": "PortalWars2.log",
}

OUT.write_text("window.SPLITSTATS_DATA = " + json.dumps(data, ensure_ascii=False, separators=(",", ":")) + ";\n", encoding="utf-8")
print(f"Wrote {OUT} ({len(data['matches'])} matches, {len(data['teams'])} team rows, {len(data['players'])} player rows, {len(data['events'])} events)")
